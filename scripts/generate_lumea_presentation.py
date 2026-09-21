import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Brand Palette (Girly, Elegant, Clean)
BG_BLUSH      = RGBColor(255, 248, 250)  # #FFF8FA
ROSE_PRIMARY  = RGBColor(217, 108, 138)  # #D96C8A
WINE_DEEP     = RGBColor(43, 32, 36)     # #2B2024
ROSE_LIGHT    = RGBColor(244, 216, 223)  # #F4D8DF
ROSE_ACCENT   = RGBColor(169, 75, 104)   # #A94B68
WHITE         = RGBColor(255, 255, 255)
CHARCOAL      = RGBColor(75, 65, 70)
CARD_BG       = RGBColor(255, 255, 255)

def apply_slide_background(slide, prs):
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_BLUSH
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title_text, category_text="LUMÉA BEAUTY PLATFORM"):
    # Category Pill
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    r_cat = p_cat.add_run()
    r_cat.text = f"✦ {category_text.upper()} ✦"
    r_cat.font.name = 'Calibri'
    r_cat.font.size = Pt(10)
    r_cat.font.bold = True
    r_cat.font.color.rgb = ROSE_PRIMARY

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.7))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.name = 'Georgia'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = WINE_DEEP

    # Bottom accent divider line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ROSE_LIGHT
    line.line.fill.background()

def add_card(slide, left, top, width, height, bg_rgb=CARD_BG, border_rgb=ROSE_LIGHT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_rgb
    card.line.color.rgb = border_rgb
    card.line.width = Pt(1.2)
    return card

def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen Dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s1, prs)

    # Decorative left accent panel
    left_banner = s1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.4), prs.slide_height
    )
    left_banner.fill.solid()
    left_banner.fill.fore_color.rgb = ROSE_PRIMARY
    left_banner.line.fill.background()

    # Content Container Card
    add_card(s1, Inches(1.0), Inches(0.8), Inches(7.2), Inches(5.9), WHITE, ROSE_LIGHT)

    # Pill tag
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(1.2), Inches(3.0), Inches(0.4))
    pill.fill.solid()
    pill.fill.fore_color.rgb = ROSE_LIGHT
    pill.line.fill.background()
    p_pill = pill.text_frame.paragraphs[0]
    p_pill.text = "BEAUTY TECH PRESENTATION"
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.runs[0].font.name = 'Calibri'
    p_pill.runs[0].font.size = Pt(9.5)
    p_pill.runs[0].font.bold = True
    p_pill.runs[0].font.color.rgb = ROSE_ACCENT

    # Main Brand Title
    tb = s1.shapes.add_textbox(Inches(1.4), Inches(1.75), Inches(6.4), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "LUMÉA"
    r1.font.name = 'Georgia'
    r1.font.size = Pt(44)
    r1.font.bold = True
    r1.font.color.rgb = WINE_DEEP

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Glow, your way"
    r2.font.name = 'Georgia'
    r2.font.size = Pt(22)
    r2.font.italic = True
    r2.font.color.rgb = ROSE_PRIMARY

    p3 = tf.add_paragraph()
    p3.space_before = Pt(10)
    r3 = p3.add_run()
    r3.text = "A Modern Single-Product Beauty E-Commerce Platform\nReact 18 · Node.js Express · MySQL · Tailwind CSS"
    r3.font.name = 'Calibri'
    r3.font.size = Pt(12)
    r3.font.color.rgb = CHARCOAL

    # Student / Project Info Box
    add_card(s1, Inches(1.4), Inches(4.3), Inches(6.4), Inches(2.0), BG_BLUSH, ROSE_LIGHT)
    tb_info = s1.shapes.add_textbox(Inches(1.6), Inches(4.4), Inches(6.0), Inches(1.8))
    tf_info = tb_info.text_frame
    tf_info.word_wrap = True
    
    pi1 = tf_info.paragraphs[0]
    r_lbl1 = pi1.add_run(); r_lbl1.text = "Project by: "; r_lbl1.font.bold = True
    r_val1 = pi1.add_run(); r_val1.text = "Kichu Khoirom (University Roll No. 2433548)\n"
    r_lbl2 = pi1.add_run(); r_lbl2.text = "Degree: "; r_lbl2.font.bold = True
    r_val2 = pi1.add_run(); r_val2.text = "Bachelor of Computer Applications (BCA)\n"
    r_lbl3 = pi1.add_run(); r_lbl3.text = "Institution: "; r_lbl3.font.bold = True
    r_val3 = pi1.add_run(); r_val3.text = "Gulzar Group of Institutions, Khanna, Punjab\n"
    r_lbl4 = pi1.add_run(); r_lbl4.text = "Industry Internship: "; r_lbl4.font.bold = True
    r_val4 = pi1.add_run(); r_val4.text = "Viro Web Solution (Web Development)"
    for run in pi1.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.color.rgb = WINE_DEEP

    # Right Image Showcase
    hero_img = "public/images/lumea_hero_tint.png"
    if os.path.exists(hero_img):
        add_card(s1, Inches(8.5), Inches(0.8), Inches(4.0), Inches(5.9), WHITE, ROSE_LIGHT)
        s1.shapes.add_picture(hero_img, Inches(8.75), Inches(1.2), width=Inches(3.5))

    # =========================================================================
    # SLIDE 2: THE PRODUCT — LUMÉA GLOW TINT
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s2, prs)
    add_header(s2, "The Product — Luméa Glow Tint", "FLAGSHIP COSMETIC SHOWCASE")

    cards_data = [
        ("🌸 Multifunctional Formula", "A silky cream tint engineered to melt seamlessly into the skin for an effortless, buildable flush on cheeks, lips, and eyelids. Replaces 3 cosmetic products in 1 pocket-sized 8g component."),
        ("🌿 Clean & Conscious Beauty", "Formulated with 100% vegan botanical waxes, cruelty-free certification, zero mineral oils, and enriched with hydrating jojoba and squalane for 12-hour continuous skin dewiness."),
        ("💎 Approachable Luxury (₹799)", "Direct-to-consumer pricing with zero retail middleman markups. Ships free across India with Cash on Delivery (COD) support and a guaranteed 4.9★ customer satisfaction rating.")
    ]

    top_y = 1.7
    for title, desc in cards_data:
        add_card(s2, Inches(0.8), Inches(top_y), Inches(7.2), Inches(1.6), WHITE, ROSE_LIGHT)
        tb_c = s2.shapes.add_textbox(Inches(1.0), Inches(top_y + 0.1), Inches(6.8), Inches(1.4))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.name = 'Georgia'
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = ROSE_ACCENT

        p_d = tf_c.add_paragraph()
        p_d.space_before = Pt(4)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.name = 'Calibri'
        r_d.font.size = Pt(10.5)
        r_d.font.color.rgb = CHARCOAL
        top_y += 1.8

    specs_img = "docs/screenshots/fig_7_2_specs_gallery.png"
    if os.path.exists(specs_img):
        add_card(s2, Inches(8.3), Inches(1.7), Inches(4.2), Inches(5.2), WHITE, ROSE_LIGHT)
        s2.shapes.add_picture(specs_img, Inches(8.5), Inches(1.9), width=Inches(3.8))

    # =========================================================================
    # SLIDE 3: 6 SIGNATURE SHADES & PALETTE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s3, prs)
    add_header(s3, "Curated Color Palette — 6 Signature Shades", "SHADE SPECTRUM & INCLUSIVITY")

    shades = [
        ("Rose Petal", "#D9828B", "Soft Rosy Pink", "Universal everyday blush for cool to neutral undertones"),
        ("Peach Bloom", "#EFA07F", "Warm Coral Peach", "Sun-kissed radiance ideal for light to medium warm skin"),
        ("Berry Kiss", "#A94B68", "Deep Rich Plum", "Bold romantic statement tint for medium to deep undertones"),
        ("Soft Coral", "#E87970", "Fresh Coral Tangerine", "Lively playful flush perfect for daytime lip and cheek glow"),
        ("Nude Glow", "#B97862", "Warm Terracotta Nude", "Subtle sculpted neutral giving natural 'no-makeup' definition"),
        ("Pink Champagne", "#E8A5B5", "Shimmer Mauve Pink", "Dewy highlighted glow with ultra-fine light-reflecting mica")
    ]

    grid = [
        (0.8, 1.7), (4.8, 1.7), (8.8, 1.7),
        (0.8, 4.3), (4.8, 4.3), (8.8, 4.3)
    ]

    for idx, (s_name, hex_code, desc_short, undertone) in enumerate(shades):
        gx, gy = grid[idx]
        add_card(s3, Inches(gx), Inches(gy), Inches(3.7), Inches(2.4), WHITE, ROSE_LIGHT)
        
        dot = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(gx + 0.3), Inches(gy + 0.3), Inches(0.65), Inches(0.65))
        h = hex_code.lstrip('#')
        rgb = RGBColor(*(int(h[i:i+2], 16) for i in (0, 2, 4)))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb
        dot.line.color.rgb = WHITE
        dot.line.width = Pt(2)

        tb_sh = s3.shapes.add_textbox(Inches(gx + 1.1), Inches(gy + 0.25), Inches(2.4), Inches(1.9))
        tf_sh = tb_sh.text_frame
        tf_sh.word_wrap = True
        
        p_name = tf_sh.paragraphs[0]
        r_nm = p_name.add_run(); r_nm.text = s_name; r_nm.font.bold = True
        r_nm.font.name = 'Georgia'; r_nm.font.size = Pt(13); r_nm.font.color.rgb = WINE_DEEP

        p_hex = tf_sh.add_paragraph()
        r_hx = p_hex.add_run(); r_hx.text = f"{hex_code} · {desc_short}"; r_hx.font.size = Pt(9.5)
        r_hx.font.color.rgb = ROSE_PRIMARY; r_hx.font.bold = True

        p_u = tf_sh.add_paragraph()
        p_u.space_before = Pt(4)
        r_u = p_u.add_run(); r_u.text = undertone; r_u.font.size = Pt(9)
        r_u.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 4: THE PROBLEM & LUMÉA SOLUTION
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s4, prs)
    add_header(s4, "Market Challenge vs. The Luméa Experience", "PROBLEM & SOLUTION ANALYSIS")

    # Left: Traditional Problem
    add_card(s4, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), WHITE, RGBColor(245, 200, 210))
    tb_p = s4.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    
    p_th = tf_p.paragraphs[0]
    r_th = p_th.add_run(); r_th.text = "❌ Traditional Cosmetic E-Commerce"; r_th.font.bold = True
    r_th.font.name = 'Georgia'; r_th.font.size = Pt(14); r_th.font.color.rgb = RGBColor(190, 40, 60)

    p_body_p = tf_p.add_paragraph()
    p_body_p.space_before = Pt(12)
    r_bp = p_body_p.add_run()
    r_bp.text = (
        "• Catalog Clutter: Overwhelming marketplaces with hundreds of unrelated products cause severe decision fatigue.\n\n"
        "• Shade Hesitation: Static studio photos don't demonstrate shade warmth or skin undertone matching.\n\n"
        "• Painful Checkout: Mandatory multi-page logins, hidden shipping fees, and broken UPI options lead to 70%+ cart drop-offs.\n\n"
        "• Post-Purchase Anxiety: No live tracking visibility once an order is confirmed."
    )
    r_bp.font.name = 'Calibri'; r_bp.font.size = Pt(11); r_bp.font.color.rgb = CHARCOAL

    # Right: Luméa Solution
    add_card(s4, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), WHITE, ROSE_PRIMARY)
    tb_s = s4.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p_sh = tf_s.paragraphs[0]
    r_sh = p_sh.add_run(); r_sh.text = "✨ The Luméa Single-Product Experience"; r_sh.font.bold = True
    r_sh.font.name = 'Georgia'; r_sh.font.size = Pt(14); r_sh.font.color.rgb = ROSE_ACCENT

    p_body_s = tf_s.add_paragraph()
    p_body_s.space_before = Pt(12)
    r_bs = p_body_s.add_run()
    r_bs.text = (
        "• Laser-Focused D2C Storytelling: One hero formula perfected across 6 universally flattering shades.\n\n"
        "• Live Swatch Synchronizer: Clicking any shade instantly switches photography, ambient glow, and inventory count.\n\n"
        "• Frictionless Checkout: 1-click bag drawer, transparent pricing (₹799 with Free Shipping), and reliable Cash on Delivery.\n\n"
        "• Total Transparency: Real-time 5-stage fulfillment stepper + self-service order cancellation & ticket desk."
    )
    r_bs.font.name = 'Calibri'; r_bs.font.size = Pt(11); r_bs.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 5: TECHNOLOGY STACK & ARCHITECTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s5, prs)
    add_header(s5, "Full-Stack Technology Architecture", "ENGINEERING FOUNDATION")

    tech_blocks = [
        ("React 18 + Vite 5.4", "Single-Page Application frontend with declarative component hierarchy, custom context hooks (Cart, Auth, Notifications), and sub-second HMR updates."),
        ("Tailwind CSS 3.4", "Utility-first design system with customized tokens (#FFF8FA, #D96C8A, #2B2024), 60 FPS mobile transforms, and zero layout drag."),
        ("Node.js & Express REST API", "Asynchronous non-blocking backend server orchestrating products, orders, support tickets, and role verification."),
        ("MySQL 8.0 + In-Memory Fallback", "Relational database in 3NF with automated zero-downtime fallback to an in-memory replica if database is unconfigured.")
    ]

    ty = 1.7
    for name, desc in tech_blocks:
        add_card(s5, Inches(0.8), Inches(ty), Inches(5.8), Inches(1.2), WHITE, ROSE_LIGHT)
        tb_t = s5.shapes.add_textbox(Inches(1.0), Inches(ty + 0.05), Inches(5.4), Inches(1.1))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p_tn = tf_t.paragraphs[0]
        r_tn = p_tn.add_run(); r_tn.text = f"⚡ {name}"; r_tn.font.bold = True
        r_tn.font.name = 'Georgia'; r_tn.font.size = Pt(12); r_tn.font.color.rgb = ROSE_ACCENT
        
        p_td = tf_t.add_paragraph()
        r_td = p_td.add_run(); r_td.text = desc; r_td.font.size = Pt(9.5)
        r_td.font.name = 'Calibri'; r_td.font.color.rgb = CHARCOAL
        ty += 1.35

    arch_img = "docs/screenshots/flow_architecture.png"
    if os.path.exists(arch_img):
        add_card(s5, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.2), WHITE, ROSE_LIGHT)
        s5.shapes.add_picture(arch_img, Inches(7.1), Inches(2.2), width=Inches(5.2))

    # =========================================================================
    # SLIDE 6: HOW IT WORKS — INTERACTIVE STOREFRONT
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s6, prs)
    add_header(s6, "Interactive Storefront & Swatch Experience", "FRONTEND INTERACTIVITY & DESIGN")

    add_card(s6, Inches(0.8), Inches(1.7), Inches(5.4), Inches(5.2), WHITE, ROSE_LIGHT)
    tb_feat = s6.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(4.8), Inches(4.8))
    tf_feat = tb_feat.text_frame
    tf_feat.word_wrap = True

    pf_head = tf_feat.paragraphs[0]
    r_pf = pf_head.add_run(); r_pf.text = "Customer Journey Highlights:"; r_pf.font.bold = True
    r_pf.font.name = 'Georgia'; r_pf.font.size = Pt(14); r_pf.font.color.rgb = WINE_DEEP

    pf_body = tf_feat.add_paragraph()
    pf_body.space_before = Pt(8)
    r_pfb = pf_body.add_run()
    r_pfb.text = (
        "1. Real-Time Swatch Switching:\n"
        "Selecting any color swatch instantly updates the product photography, switches the ambient background glow, and renders stock availability.\n\n"
        "2. Formula & Benefits Breakdown:\n"
        "Dedicated sections showcase clean botanical ingredients (jojoba, squalane), cruelty-free pledge, and customer testimonials.\n\n"
        "3. Persistent Cart Drawer:\n"
        "An off-canvas bag slides into view with live quantity controls, subtotal calculations, and free delivery indicators across India."
    )
    r_pfb.font.name = 'Calibri'; r_pfb.font.size = Pt(10.5); r_pfb.font.color.rgb = CHARCOAL

    hero_snap = "docs/screenshots/fig_7_1_storefront_hero.png"
    if os.path.exists(hero_snap):
        add_card(s6, Inches(6.5), Inches(1.7), Inches(6.0), Inches(5.2), WHITE, ROSE_LIGHT)
        s6.shapes.add_picture(hero_snap, Inches(6.7), Inches(1.9), width=Inches(5.6))

    # =========================================================================
    # SLIDE 7: FRICTIONLESS CHECKOUT & 5-STAGE TRACKING
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s7, prs)
    add_header(s7, "Checkout, Fulfillment & Order Tracking", "TRANSACTIONAL INTEGRITY")

    add_card(s7, Inches(0.8), Inches(1.7), Inches(5.4), Inches(2.5), WHITE, ROSE_LIGHT)
    tb_c1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.0), Inches(2.2))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    r_c1_h = tf_c1.paragraphs[0].add_run(); r_c1_h.text = "🚚 Frictionless Checkout Flow"; r_c1_h.font.bold = True
    r_c1_h.font.name = 'Georgia'; r_c1_h.font.size = Pt(12); r_c1_h.font.color.rgb = ROSE_ACCENT
    p_c1_b = tf_c1.add_paragraph()
    p_c1_b.space_before = Pt(4)
    r_c1_b = p_c1_b.add_run()
    r_c1_b.text = (
        "• Validates name, phone, city, and shipping address.\n"
        "• Zero surprise fees: transparent ₹799 subtotal with free shipping.\n"
        "• Active Cash on Delivery (COD) + friendly maintenance banner for UPI online payments."
    )
    r_c1_b.font.size = Pt(9.5); r_c1_b.font.name = 'Calibri'; r_c1_b.font.color.rgb = CHARCOAL

    add_card(s7, Inches(0.8), Inches(4.4), Inches(5.4), Inches(2.5), WHITE, ROSE_LIGHT)
    tb_c2 = s7.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(5.0), Inches(2.2))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    r_c2_h = tf_c2.paragraphs[0].add_run(); r_c2_h.text = "📦 5-Stage Visual Order Stepper"; r_c2_h.font.bold = True
    r_c2_h.font.name = 'Georgia'; r_c2_h.font.size = Pt(12); r_c2_h.font.color.rgb = ROSE_ACCENT
    p_c2_b = tf_c2.add_paragraph()
    p_c2_b.space_before = Pt(4)
    r_c2_b = p_c2_b.add_run()
    r_c2_b.text = (
        "• Search by order code (e.g. #LM1024) to inspect live progress.\n"
        "• Visual Stepper: Pending ➔ Confirmed ➔ Processing ➔ Shipped ➔ Delivered.\n"
        "• Customer Self-Cancellation: Allowed during pending state; restores inventory units automatically."
    )
    r_c2_b.font.size = Pt(9.5); r_c2_b.font.name = 'Calibri'; r_c2_b.font.color.rgb = CHARCOAL

    track_snap = "docs/screenshots/fig_7_6_track_order.png"
    if os.path.exists(track_snap):
        add_card(s7, Inches(6.5), Inches(1.7), Inches(6.0), Inches(5.2), WHITE, ROSE_LIGHT)
        s7.shapes.add_picture(track_snap, Inches(6.7), Inches(2.0), width=Inches(5.6))

    # =========================================================================
    # SLIDE 8: ADMINISTRATIVE CONTROL & RBAC SECURITY
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s8, prs)
    add_header(s8, "Administrative Control & RBAC Governance", "ROLE-BASED SECURITY")

    add_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), WHITE, ROSE_LIGHT)
    tb_rbac = s8.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(5.2), Inches(4.8))
    tf_rbac = tb_rbac.text_frame
    tf_rbac.word_wrap = True
    
    r_rb_h = tf_rbac.paragraphs[0].add_run(); r_rb_h.text = "Role-Based Access Hierarchy:"; r_rb_h.font.bold = True
    r_rb_h.font.name = 'Georgia'; r_rb_h.font.size = Pt(13); r_rb_h.font.color.rgb = WINE_DEEP

    p_rb_body = tf_rbac.add_paragraph()
    p_rb_body.space_before = Pt(8)
    r_rbb = p_rb_body.add_run()
    r_rbb.text = (
        "1. Super Admin (kichuKhoirom@gmail.com):\n"
        "• Full administrative rights.\n"
        "• Promote customers to Sub-Admin, demote roles, delete users.\n"
        "• Access real-time revenue KPIs and support tickets.\n\n"
        "2. Sub-Admin (pardeepsign@gmail.com):\n"
        "• Order status management (update to Confirmed/Shipped).\n"
        "• Customer support ticket resolution.\n"
        "• Product inventory inspection.\n\n"
        "3. Customer (james@gmail.com):\n"
        "• Storefront shopping, order tracking, and ticket submission.\n"
        "• Access to /admin is strictly blocked by Route Guard."
    )
    r_rbb.font.name = 'Calibri'; r_rbb.font.size = Pt(9.5); r_rbb.font.color.rgb = CHARCOAL

    admin_snap = "docs/screenshots/fig_7_8_admin_dashboard.png"
    if os.path.exists(admin_snap):
        add_card(s8, Inches(6.7), Inches(1.7), Inches(5.8), Inches(5.2), WHITE, ROSE_LIGHT)
        s8.shapes.add_picture(admin_snap, Inches(6.9), Inches(2.0), width=Inches(5.4))

    # =========================================================================
    # SLIDE 9: PERFORMANCE & RESILIENCE ARCHITECTURE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s9, prs)
    add_header(s9, "Performance Profiling & Fault Tolerance", "STABILITY & RESILIENCE")

    resilience_cards = [
        ("⚡ 60 FPS Mobile Optimization", "Heavyweight CSS backdrop-filter blurs were eliminated from general cards. On screens <= 640px, blurs are disabled entirely; on desktop, blur is isolated to the sticky header and hardware-accelerated with translateZ(0). Zero jank or drag."),
        ("🛡️ Multi-Tier Error Containment", "Custom 404 Page renders a luxury gradient numeral with store recovery CTAs; dedicated /503 page offers retry connectivity; root-level React ErrorBoundary prevents blank-screen whiteouts."),
        ("🔄 Zero-Downtime Database Fallback", "If the local MySQL database is offline or unconfigured, the backend automatically transitions to an in-memory replica pre-loaded with products, variants, sample orders, and user roles.")
    ]

    ty9 = 1.7
    for title, desc in resilience_cards:
        add_card(s9, Inches(0.8), Inches(ty9), Inches(7.4), Inches(1.6), WHITE, ROSE_LIGHT)
        tb_rc = s9.shapes.add_textbox(Inches(1.0), Inches(ty9 + 0.1), Inches(7.0), Inches(1.4))
        tf_rc = tb_rc.text_frame
        tf_rc.word_wrap = True
        p_rct = tf_rc.paragraphs[0]
        r_rct = p_rct.add_run(); r_rct.text = title; r_rct.font.bold = True
        r_rct.font.name = 'Georgia'; r_rct.font.size = Pt(13); r_rct.font.color.rgb = ROSE_ACCENT

        p_rcd = tf_rc.add_paragraph()
        p_rcd.space_before = Pt(4)
        r_rcd = p_rcd.add_run(); r_rcd.text = desc; r_rcd.font.size = Pt(10)
        r_rcd.font.name = 'Calibri'; r_rcd.font.color.rgb = CHARCOAL
        ty9 += 1.8

    mobile_snap = "docs/screenshots/fig_7_11_mobile_view.png"
    if os.path.exists(mobile_snap):
        add_card(s9, Inches(8.5), Inches(1.7), Inches(4.0), Inches(5.2), WHITE, ROSE_LIGHT)
        s9.shapes.add_picture(mobile_snap, Inches(9.2), Inches(1.9), width=Inches(2.6))

    # =========================================================================
    # SLIDE 10: FUTURE ROADMAP & CONCLUSION
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_slide_background(s10, prs)
    add_header(s10, "Future Roadmap & Project Conclusion", "WHAT'S NEXT FOR LUMÉA")

    roadmap_items = [
        ("💳 Live Payment Gateway Integration", "Direct integration with Razorpay and Cashfree webhooks for instant UPI QR code generation, credit/debit cards, and automated payment reconciliation."),
        ("🤳 AI-Powered Webcam Shade Matcher", "Client-side computer vision engine using TensorFlow.js face landmark detection to analyze skin undertones in natural lighting and recommend the ideal tint shade."),
        ("📲 Automated WhatsApp & SMS Alerts", "Instant order dispatch notifications, live delivery agent tracking links, and review requests powered by Gupshup / Twilio webhook messaging."),
        ("🌐 Regional Language Localization", "Multilingual store localization (Hindi, Punjabi, English) to dramatically expand reach across diverse beauty customer segments in India.")
    ]

    coords = [(0.8, 1.7), (6.8, 1.7), (0.8, 4.3), (6.8, 4.3)]
    for idx, (title, desc) in enumerate(roadmap_items):
        rx, ry = coords[idx]
        add_card(s10, Inches(rx), Inches(ry), Inches(5.7), Inches(2.4), WHITE, ROSE_LIGHT)
        tb_rm = s10.shapes.add_textbox(Inches(rx + 0.2), Inches(ry + 0.15), Inches(5.3), Inches(2.0))
        tf_rm = tb_rm.text_frame
        tf_rm.word_wrap = True

        p_rmt = tf_rm.paragraphs[0]
        r_rmt = p_rmt.add_run(); r_rmt.text = title; r_rmt.font.bold = True
        r_rmt.font.name = 'Georgia'; r_rmt.font.size = Pt(12); r_rmt.font.color.rgb = ROSE_ACCENT

        p_rmd = tf_rm.add_paragraph()
        p_rmd.space_before = Pt(6)
        r_rmd = p_rmd.add_run(); r_rmd.text = desc; r_rmd.font.size = Pt(9.5)
        r_rmd.font.name = 'Calibri'; r_rmd.font.color.rgb = CHARCOAL

    output_pptx = "docs/LUMEA_Beauty_Product_Presentation.pptx"
    prs.save(output_pptx)
    print(f"PRESENTATION SAVED: {output_pptx} ({os.path.getsize(output_pptx)} bytes)")

if __name__ == '__main__':
    build_presentation()
